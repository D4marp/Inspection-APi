from pathlib import Path
import re
import subprocess
import tempfile
import unicodedata

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from pypdf import PdfReader, PdfWriter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Ebook_Flutter_Android_Profesional_backup_tanpa_screenshot.docx"
OUTPUT_DOCX = ROOT / "Ebook_Flutter_Android_Profesional.docx"
OUTPUT_PDF = ROOT / "Ebook_Flutter_Android_Profesional.pdf"
AUTHOR = "Damar Galih Aji Pradana, S.Tr.Kom."
COVER_PNG = ROOT / "assets/cover/cover-v2.png"
COVER_SVG = ROOT / "assets/cover/cover-v2.svg"
COVER_PDF = ROOT / "assets/cover/cover-v2.pdf"

SCREENSHOTS = [
    (
        "Hari 1 — UI, daftar, dan state lokal",
        ROOT / "flutter_projects/hari_1_dpl_inspection/screenshots/hari_1_daftar.png",
        "Daftar inspeksi berjalan pada iPhone 16 Simulator.",
    ),
    (
        "Hari 2 — REST API, session, dan cache",
        ROOT / "flutter_projects/hari_2_dpl_inspection_api/screenshots/hari_2_api.png",
        "Data mock API berhasil dimuat setelah login otomatis khusus dokumentasi.",
    ),
    (
        "Hari 3 — detail dan lampiran foto",
        ROOT / "flutter_projects/hari_3_dpl_inspection_release/screenshots/hari_3_foto.png",
        "Panel Galeri/Kamera pada detail inspeksi; kamera divalidasi pada perangkat fisik.",
    ),
]


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def insert_paragraph_after(paragraph: Paragraph, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    if style:
        inserted.style = style
    return inserted


def normalized(text: str) -> str:
    value = unicodedata.normalize("NFKC", text).replace("\u00a0", " ")
    return re.sub(r"\s+", " ", value).strip().lower()


def export_pdf(docx_path: Path, output_dir: Path) -> Path:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    return output_dir / f"{docx_path.stem}.pdf"


def render_cover() -> None:
    subprocess.run(
        [
            "rsvg-convert",
            "-w",
            "1748",
            "-h",
            "2480",
            str(COVER_SVG),
            "-o",
            str(COVER_PNG),
        ],
        check=True,
    )
    subprocess.run(
        [
            "rsvg-convert",
            "-f",
            "pdf",
            "--page-width",
            "148mm",
            "--page-height",
            "210mm",
            "--width",
            "148mm",
            "--height",
            "210mm",
            str(COVER_SVG),
            "-o",
            str(COVER_PDF),
        ],
        check=True,
    )


def replace_pdf_cover(book_pdf: Path) -> None:
    book = PdfReader(str(book_pdf))
    cover = PdfReader(str(COVER_PDF))
    writer = PdfWriter()
    writer.add_page(cover.pages[0])
    for page in book.pages[1:]:
        writer.add_page(page)
    writer.add_metadata({
        "/Title": "Flutter Android & iOS — Dari UI hingga Release",
        "/Author": AUTHOR,
        "/Subject": "Buku pelatihan Flutter profesional program 3 hari",
    })
    temporary = book_pdf.with_suffix(".cover.pdf")
    with temporary.open("wb") as output:
        writer.write(output)
    temporary.replace(book_pdf)


def page_texts(pdf_path: Path) -> list[str]:
    result = subprocess.run(
        ["pdftotext", "-layout", str(pdf_path), "-"],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return result.stdout.decode("utf-8", errors="replace").split("\f")


def find_heading_page(title: str, pages: list[str]) -> int:
    needle = normalized(title)
    exact_matches: list[int] = []
    fallback_matches: list[int] = []
    for page_number, page in enumerate(pages, start=1):
        lines = [normalized(line) for line in page.splitlines()]
        if needle in lines:
            exact_matches.append(page_number)
        elif needle and needle in normalized(page):
            fallback_matches.append(page_number)
    matches = exact_matches or fallback_matches
    if not matches:
        raise RuntimeError(f"Judul tidak ditemukan pada PDF: {title}")
    return max(matches)


def build_document() -> tuple[Document, list[tuple[str, Paragraph]]]:
    document = Document(SOURCE)
    document.core_properties.author = AUTHOR
    document.core_properties.last_modified_by = AUTHOR

    # The original cover logos live in the first, otherwise-empty table.
    # Remove that table so the redesigned cover does not duplicate them.
    original_cover_table = document.tables[0]._element
    original_cover_table.getparent().remove(original_cover_table)

    paragraphs = list(document.paragraphs)
    cover_page_break = paragraphs[4]
    for paragraph in paragraphs[:4]:
        delete_paragraph(paragraph)
    cover_paragraph = cover_page_break
    cover_paragraph.clear()
    cover_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_paragraph.paragraph_format.space_before = Pt(0)
    cover_paragraph.paragraph_format.space_after = Pt(0)
    usable_width = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
    # Keep enough room for Word/LibreOffice to place the following page-break
    # heading on page 1 before advancing to the TOC. The PDF cover is replaced
    # separately with the full-bleed A5 artwork.
    cover_paragraph.add_run().add_picture(str(COVER_PNG), width=int(usable_width * 0.80))
    document.sections[0].different_first_page_header_footer = True

    document.add_page_break()
    document.add_heading("Bukti Aplikasi pada iPhone 16 Simulator", level=1)
    document.add_paragraph(
        "Seluruh gambar pada lampiran ini diambil langsung dari iPhone 16 Simulator "
        "iOS 18.0, bukan mockup. Toolchain verifikasi: Flutter 3.44.8 stable, "
        "Dart 3.12.2, DevTools 2.57.0, dan Xcode 26.5."
    )
    for title, image_path, caption in SCREENSHOTS:
        document.add_heading(title, level=2)
        picture = document.add_picture(str(image_path), width=Inches(3.0))
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph = document.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style
        and paragraph.style.name == "Heading 1"
        and paragraph.text.strip() != "Daftar isi"
    ]
    toc_heading = next(
        paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Daftar isi"
    )
    toc_heading.paragraph_format.page_break_before = True
    next_heading = headings[0]
    paragraphs = list(document.paragraphs)
    start = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph._p is toc_heading._p
    ) + 1
    end = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph._p is next_heading._p
    )
    for paragraph in paragraphs[start:end]:
        delete_paragraph(paragraph)

    style_name = "Daftar Isi Statis"
    if style_name not in document.styles:
        style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(8)

    section = document.sections[0]
    tab_position = section.page_width - section.left_margin - section.right_margin
    anchor = toc_heading
    entries: list[tuple[str, Paragraph]] = []
    for heading in headings:
        entry = insert_paragraph_after(anchor, style_name)
        entry.paragraph_format.space_before = Pt(0)
        entry.paragraph_format.space_after = Pt(0)
        entry.paragraph_format.line_spacing = 1
        entry.paragraph_format.tab_stops.add_tab_stop(
            tab_position, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        entry.add_run(heading.text.strip())
        entry.add_run("\t")
        entry.add_run("000")
        entries.append((heading.text.strip(), entry))
        anchor = entry

    page_break = insert_paragraph_after(anchor)
    page_break.add_run().add_break(WD_BREAK.PAGE)
    return document, entries


def main() -> None:
    missing = [str(path) for _, path, _ in SCREENSHOTS if not path.exists()]
    if missing:
        raise SystemExit("Screenshot tidak ditemukan:\n" + "\n".join(missing))

    render_cover()
    document, entries = build_document()
    with tempfile.TemporaryDirectory() as temporary:
        temp_dir = Path(temporary)
        draft_docx = temp_dir / "draft.docx"
        document.save(draft_docx)
        draft_pdf = export_pdf(draft_docx, temp_dir)
        pages = page_texts(draft_pdf)

        for title, paragraph in entries:
            paragraph.runs[-1].text = str(find_heading_page(title, pages))

        document.save(OUTPUT_DOCX)

    export_pdf(OUTPUT_DOCX, ROOT)
    replace_pdf_cover(OUTPUT_PDF)
    final_pages = page_texts(OUTPUT_PDF)
    if len(entries) < 100:
        raise RuntimeError("Daftar isi tidak lengkap")
    if "01. flutter sebagai ui deklaratif" not in normalized("\n".join(final_pages[1:8])):
        raise RuntimeError("Daftar isi tidak tercetak pada halaman awal")

    print(f"Ebook dipulihkan dari sumber asli: {OUTPUT_DOCX}")
    print(f"Daftar isi statis: {len(entries)} entri")
    print(f"PDF final: {OUTPUT_PDF} ({len(final_pages) - 1} halaman)")


if __name__ == "__main__":
    main()
